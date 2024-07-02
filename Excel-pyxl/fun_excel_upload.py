import openpyxl


def to_update_excel(file_path,fruit_name,changed_price):
    book = openpyxl.load_workbook(file_path)
    sheet = book.active
    fruit_coll = 0
    p_coll = 0
    fruit_row = 0
    # to find the column use price name
    for i in range(1, sheet.max_column + 1):
        if sheet.cell(row=1, column=i).value == "price":
            p_coll = i
        if sheet.cell(row=1, column=i).value == "fruit_name":
            fruit_coll = i

    # to find the name of fruit
    for i in range(1, sheet.max_row + 1):
        if sheet.cell(row=i, column=fruit_coll).value == fruit_name:
            fruit_row = i

    sheet.cell(row=fruit_row, column=p_coll).value = changed_price
    book.save(file_path)